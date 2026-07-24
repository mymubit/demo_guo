# -*- coding: utf-8 -*-
"""将 episode_scripts 产物导出为 Word (.docx)。"""
from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document


def build_episode_scripts_docx(*, project_title: str, scripts_payload: dict) -> bytes:
    """返回 docx 文件字节；含标题与各集正文。"""
    document = Document()
    title = (project_title or "").strip() or "未命名项目"
    document.add_heading(title, level=0)

    episodes = scripts_payload.get("episodes") if isinstance(scripts_payload, dict) else None
    if not isinstance(episodes, list):
        episodes = []

    for episode in episodes:
        if not isinstance(episode, dict):
            continue
        number = episode.get("episode_number")
        ep_title = str(episode.get("title") or "").strip()
        heading = _episode_heading(number=number, title=ep_title)
        document.add_heading(heading, level=1)
        body = _episode_body(episode)
        if body:
            for paragraph in body.split("\n"):
                document.add_paragraph(paragraph)
        else:
            document.add_paragraph("")

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _episode_heading(*, number: Any, title: str) -> str:
    if number is None or number == "":
        return title or "未命名集"
    try:
        ep_no = int(number)
    except (TypeError, ValueError):
        ep_no = number
    if title:
        return f"第 {ep_no} 集 {title}"
    return f"第 {ep_no} 集"


def _episode_body(episode: dict) -> str:
    script = episode.get("script")
    if isinstance(script, str) and script.strip():
        return script.strip()
    return _scenes_to_text(episode.get("scenes"))


def _scenes_to_text(scenes: Any) -> str:
    if not isinstance(scenes, list):
        return ""
    lines: list[str] = []
    for scene in scenes:
        if not isinstance(scene, dict):
            continue
        heading = str(scene.get("heading") or "").strip()
        if heading:
            lines.append(heading)
        for beat in scene.get("beats") or []:
            if not isinstance(beat, dict):
                continue
            text = str(beat.get("text") or "").strip()
            if beat.get("type") == "dialogue":
                character = str(beat.get("character") or "").strip()
                lines.append(f"{character}: {text}" if character else text)
            elif text:
                lines.append(text)
    return "\n".join(line for line in lines if line).strip()
