# -*- coding: utf-8 -*-
"""V3 P3-W2：docx 生成服务 + Delivery 导出 API。"""
from __future__ import annotations

import copy
import json
from io import BytesIO
from pathlib import Path

from django.contrib.auth import get_user_model
from django.test import SimpleTestCase, override_settings
from docx import Document
from rest_framework.test import APITestCase

from apps.drama.models import V3ArtifactVersion, V3Project
from apps.drama.orchestrator.docx_export import build_episode_scripts_docx
from apps.drama.orchestrator.report_meta import attach_script_meta
from apps.drama.skills_bridge.recipe_map import recipe_for
from apps.drama.tests.helpers import SKILLS_ROOT

_FIXTURE_DIR = Path(__file__).resolve().parent / "fixtures"
_BRIEF = json.loads(
    (_FIXTURE_DIR / "v3_project_brief_candidate.json").read_text(encoding="utf-8")
)
_BLUEPRINT = json.loads(
    (_FIXTURE_DIR / "v3_blueprint_bundle.json").read_text(encoding="utf-8")
)
_EPISODE_PLAN = json.loads(
    (_FIXTURE_DIR / "v3_episode_plan_candidate.json").read_text(encoding="utf-8")
)
_SCRIPTS = json.loads(
    (_FIXTURE_DIR / "v3_episode_scripts_candidate.json").read_text(encoding="utf-8")
)
_QUALITY = json.loads(
    (_FIXTURE_DIR / "v3_quality_report.json").read_text(encoding="utf-8")
)
_COMPLIANCE = json.loads(
    (_FIXTURE_DIR / "v3_compliance_report.json").read_text(encoding="utf-8")
)
_BLUEPRINT_KEYS = recipe_for("generate_blueprint")["writes"]
_DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


class BuildEpisodeScriptsDocxTests(SimpleTestCase):
    def test_bytes_are_zip_docx_and_contain_title(self) -> None:
        payload = {
            "episodes": [
                {
                    "episode_number": 1,
                    "title": "重生归来",
                    "script": "林夏查看授权书，发现日期异常。",
                }
            ]
        }
        raw = build_episode_scripts_docx(
            project_title="导出单元测试",
            scripts_payload=payload,
        )
        self.assertTrue(raw.startswith(b"PK\x03\x04"))
        doc = Document(BytesIO(raw))
        texts = [p.text for p in doc.paragraphs if p.text.strip()]
        self.assertIn("导出单元测试", texts)
        self.assertTrue(any("第 1 集" in t and "重生归来" in t for t in texts))
        self.assertIn("林夏查看授权书，发现日期异常。", texts)

    def test_scenes_fallback_when_script_missing(self) -> None:
        payload = {
            "episodes": [
                {
                    "episode_number": 2,
                    "title": "夜谈",
                    "scenes": [
                        {
                            "heading": "2-1 NIGHT INT. 书房",
                            "beats": [
                                {"type": "action", "text": "林夏关上门。"},
                                {
                                    "type": "dialogue",
                                    "character": "林夏",
                                    "text": "今晚必须说清楚。",
                                },
                            ],
                        }
                    ],
                }
            ]
        }
        raw = build_episode_scripts_docx(
            project_title="场景导出",
            scripts_payload=payload,
        )
        doc = Document(BytesIO(raw))
        texts = [p.text for p in doc.paragraphs if p.text.strip()]
        self.assertIn("2-1 NIGHT INT. 书房", texts)
        self.assertIn("林夏关上门。", texts)
        self.assertIn("林夏: 今晚必须说清楚。", texts)


@override_settings(DRAMA_SKILLS_ROOT=SKILLS_ROOT)
class V3DeliveryExportDocxApiTests(APITestCase):
    def setUp(self) -> None:
        user_model = get_user_model()
        self.user = user_model.objects.create_user(
            username="docx_export_api", password="pass12345"
        )
        self.other = user_model.objects.create_user(
            username="docx_export_other", password="pass12345"
        )
        self.client.force_authenticate(user=self.user)
        self.project = V3Project.objects.create(
            owner=self.user,
            title="Docx 导出项目",
            entry_type=V3Project.EntryType.ORIGINAL,
            stage=V3Project.Stage.QUALITY,
        )
        self.url = f"/api/v3/projects/{self.project.id}/delivery/export/docx/"

    def _commit(self, key: str, payload: dict, *, version: int = 1) -> V3ArtifactVersion:
        return V3ArtifactVersion.objects.create(
            project=self.project,
            artifact_key=key,
            version=version,
            status=V3ArtifactVersion.Status.COMMITTED,
            payload=payload,
        )

    def _seed_writing_deps(self) -> V3ArtifactVersion:
        self._commit("project_brief", _BRIEF)
        for key in _BLUEPRINT_KEYS:
            self._commit(key, _BLUEPRINT[key])
        self._commit("episode_plan", _EPISODE_PLAN)
        return self._commit("episode_scripts", copy.deepcopy(_SCRIPTS))

    def _seed_passing_reports(self, scripts: V3ArtifactVersion) -> None:
        self._commit(
            "quality_report",
            attach_script_meta(copy.deepcopy(_QUALITY), script_art=scripts),
        )
        self._commit(
            "compliance_report",
            attach_script_meta(copy.deepcopy(_COMPLIANCE), script_art=scripts),
        )

    def test_export_gate_fail_returns_400_with_blockers(self) -> None:
        self._seed_writing_deps()
        resp = self.client.post(self.url, {}, format="json")
        self.assertEqual(resp.status_code, 400)
        body = resp.json()
        self.assertEqual(body["code"], 400)
        self.assertIn("门禁", body["message"])
        data = body["data"]
        self.assertFalse(data["passed"])
        self.assertTrue(data["blockers"])
        self.assertTrue(any("质量报告" in b for b in data["blockers"]))

    def test_export_success_returns_docx_attachment(self) -> None:
        scripts = self._seed_writing_deps()
        self._seed_passing_reports(scripts)

        resp = self.client.post(self.url, {}, format="json")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp["Content-Type"].split(";")[0], _DOCX_CONTENT_TYPE)
        disposition = resp["Content-Disposition"]
        self.assertIn("attachment", disposition)
        self.assertIn(".docx", disposition)

        raw = b"".join(resp.streaming_content)
        self.assertTrue(raw.startswith(b"PK\x03\x04"))
        doc = Document(BytesIO(raw))
        texts = [p.text for p in doc.paragraphs if p.text.strip()]
        self.assertIn("Docx 导出项目", texts)
        self.assertTrue(any("重生归来" in t for t in texts))

    def test_owner_isolation_returns_404(self) -> None:
        foreign = V3Project.objects.create(
            owner=self.other,
            title="他人导出",
            entry_type=V3Project.EntryType.ORIGINAL,
        )
        resp = self.client.post(
            f"/api/v3/projects/{foreign.id}/delivery/export/docx/",
            {},
            format="json",
        )
        self.assertEqual(resp.status_code, 404)
